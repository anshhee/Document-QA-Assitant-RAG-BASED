from docx import Document as DocxDocument 
from pathlib import Path
from uuid import uuid4
from src.ingestion.models import Document, DocumentSection 

def load_docx(
    file_path: str | Path,
    original_filename: str | None = None,
) -> Document:
    """
    Load a DOCX file and group paragraphs by heading.
    DOCX page numbers are not preserved because they are not
    reliably available from python-docx.
    """
    file_path = Path(file_path)

    if not file_path.exists():
        raise FileNotFoundError(f"DOCX file not found: {file_path}")

    if file_path.suffix.lower() != ".docx":
        raise ValueError(f"Expected a DOCX file, got: {file_path.suffix}")

    filename = original_filename or file_path.name
    docx = DocxDocument(str(file_path))

    sections: list[DocumentSection] = []
    current_heading = None
    current_paragraphs: list[str] = []

    def save_section():
        text = "\n".join(current_paragraphs).strip()

        if text:
            sections.append(
                DocumentSection(
                    heading=current_heading,
                    text=text,
                    page=None,
                )
            )

    for paragraph in docx.paragraphs:
        text = paragraph.text.strip()

        if not text:
            continue

        style_name = paragraph.style.name if paragraph.style else ""

        if style_name.startswith("Heading"):
            save_section()
            current_heading = text
            current_paragraphs = []
        else:
            current_paragraphs.append(text)

    save_section()

    return Document(
        document_id=str(uuid4()),
        title=Path(filename).stem,
        filename=filename,
        source=filename,
        file_type="docx",
        sections=sections,
    )
